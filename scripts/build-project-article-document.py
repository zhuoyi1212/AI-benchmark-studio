"""Create the project and article collection from the sanitized DB snapshot."""

from datetime import datetime
import importlib.util
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "docs/project-articles"
SOURCE = FOLDER / "项目与文章汇编数据.json"
OUTPUT = FOLDER / "已测评项目与文章汇编.docx"
spec = importlib.util.spec_from_file_location("portfolio_style", ROOT / "scripts/build-portfolio-document.py")
style = importlib.util.module_from_spec(spec)
spec.loader.exec_module(style)


def local_time(value):
    if not value:
        return "未记录"
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")


def build():
    source = json.loads(SOURCE.read_text())
    totals = source["totals"]
    doc = Document()
    md = []
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.67)
    for name, size in [("Normal", 10.5), ("Title", 22), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        s = doc.styles[name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string("000000")
        s.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.2
        if name.startswith("Heading"):
            s.font.bold = True
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.keep_with_next = True
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.core_properties.title = "已测评项目与文章汇编"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style.font(footer.add_run("AI Benchmark Studio   "), 8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def paragraph(text, bold=False, size=None, markdown=True, keep_next=False):
        for line in str(text).split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = keep_next
            if line:
                style.font(p.add_run(line), size, bold)
            else:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = Pt(3)
        if markdown:
            md.extend([str(text), ""])

    def heading(text, level=1):
        p = doc.add_paragraph(style="Title" if level == 0 else f"Heading {level}")
        style.inline(p, text)
        md.extend(["#" * (level + 1) + " " + text, ""])

    def table(headers, rows, widths):
        # Convert only structural tables; article body text is always preserved verbatim.
        lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
        lines.extend("| " + " | ".join(str(v).replace("\n", " / ").replace("|", "｜") for v in row) + " |" for row in rows)
        style.add_table(doc, lines)
        actual = doc.tables[-1]
        for column, width in zip(actual.columns, widths):
            column.width = Inches(width)
        for row in actual.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
        # These compact indexes and evidence tables fit a page as a unit.
        for row in actual.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        md.extend(lines + [""])

    heading("已测评项目与文章汇编", 0)
    paragraph(f"数据截至 {local_time(source['capturedAt'])} 北京时间", size=10)
    paragraph(f"共{totals['projects']}个项目、{totals['articleVersions']}个已保存文章版本，包含{totals['manualVersions']}个人工编辑版本。下文按项目汇总测评信息，先展示最新已就绪稿，再按版本倒序收录历史全文。")
    paragraph("范围为App当前可见、有原始实测证据或文章正文的真实项目，不含内部验证和已删除项目。建项意图、任务提示词及文章均保留原文；ready表示已保存，不等于已经发表或人工采用。历史稿未重新核验事实，数据与价格以原文时点为准。", size=10)
    table(["项目", "最新已就绪文章", "版本数", "证据数"], [[p["name"], next((a["title"] for a in p["articles"] if a["isLatestReady"]), "暂无ready稿"), len(p["articles"]), p["artifactCount"]] for p in source["projects"]], [1.65, 4.08, 0.60, 0.60])
    paragraph("证据数包含该项目历史方案的已上传文件；正文配图仅列该文章版本实际保存的图片关联。Word中保留文件名、说明和对应段落，原始图片、PDF及视频继续保存在App，不在本汇编中复制二进制文件。", size=10)

    for number, project in enumerate(source["projects"], 1):
        doc.add_page_break()
        md.extend(["<!-- page -->", ""])
        heading(f"项目{number:02}")
        paragraph(project["name"], bold=True, size=15)
        products = "、".join(p["name"] for p in project["products"])
        mode = "快速模式" if project["mode"] == "quick" else "专业模式"
        paragraph(f"产品 {products}   {mode}   原始证据 {project['artifactCount']}份   文章 {len(project['articles'])}版", size=10)
        paragraph("在线项目 " + project["url"], size=9)
        heading("建项意图原文", 2)
        paragraph(project["intent"])
        protocol = project["protocol"]
        if protocol:
            heading("当前方案与测试任务", 2)
            paragraph(f"方案v{protocol['version']}  {protocol['title']}  {'已锁定' if protocol['status'] == 'locked' else '草稿'}", size=10)
            for case in protocol["cases"]:
                paragraph(case["title"], bold=True)
                paragraph("任务目标 " + case["objective"])
                paragraph("原始提示词", bold=True, size=10)
                paragraph(case["prompt"])
        findings = project["findings"]
        if findings:
            heading("已确认测评发现", 2)
            paragraph(f"发现v{findings['version']}  确认于{local_time(findings['confirmedAt'])}", size=9)
            paragraph(findings["summary"])
            if findings["limitations"]:
                paragraph("记录中的限制", bold=True, size=10)
                limitations = findings["limitations"]
                for item in limitations if isinstance(limitations, list) else [limitations]:
                    paragraph("• " + str(item), size=10)
        else:
            paragraph("该项目未保存人工确认的发现摘要。", size=10)
        heading("文章版本索引", 2)
        table(["版本", "保存时间", "类型", "文章标题"], [[f"v{a['version']}", local_time(a["createdAt"]), "人工编辑" if a["manualEdit"] else "保存稿", a["title"]] for a in project["articles"]], [0.48, 1.30, 0.80, 4.35])
        seen = {}
        for article in project["articles"]:
            heading(("最新已就绪文章" if article["isLatestReady"] else "历史文章") + f" v{article['version']}", 2)
            if article["contentSha256"] in seen:
                paragraph(f"标题、正文、标签与互动句均与v{seen[article['contentSha256']]}完全相同，全文见该版本。", size=10)
                continue
            seen[article["contentSha256"]] = article["version"]
            paragraph(article["title"], bold=True, size=13, keep_next=True)
            kind = "人工编辑稿，不标记为AI已审" if article["manualEdit"] and not article["aiReviewed"] else "已保存版本"
            paragraph(f"版本v{article['version']}   保存于{local_time(article['createdAt'])}   状态{article['status']}   {kind}", size=9, keep_next=True)
            paragraph(article["body"])
            tags = article["tags"]
            paragraph("标签 " + (" / ".join(tags) if isinstance(tags, list) else str(tags)), size=10)
            paragraph("互动句 " + (article["interactionQuestion"] or "未填写"), size=10)
            paragraph(f"记录模型 {article['provider'] or '未记录'} / {article['model'] or '未记录'}；记录ID {article['id']}", size=8)
            if article["manualEdit"]:
                paragraph("模型字段可能沿用来源版本，不代表此次人工修改重新调用了模型。", size=9)
            if article["evidence"]:
                heading("本版本配图对应关系", 3)
                table(["原始文件与位置", "对应说明"], [[f"第{item['order']}张 {item['fileName']}" + (f" 正文第{item['supportsParagraph']}段" if item["supportsParagraph"] else ""), item["reason"] + (" 支持原句 " + item["supportsClaim"] if item["supportsClaim"] else "")] for item in article["evidence"]], [2.10, 4.83])
            else:
                paragraph("该文章版本未保存配图关联。", size=9)
    doc.save(OUTPUT)
    # Normalize Markdown line endings only; JSON keeps the exact source strings.
    (FOLDER / "已测评项目与文章汇编.md").write_text("\n".join(md).replace("\r\n", "\n").replace("\r", "\n"))
    print(json.dumps({"document": str(OUTPUT), **totals}, ensure_ascii=False))


if __name__ == "__main__":
    build()
