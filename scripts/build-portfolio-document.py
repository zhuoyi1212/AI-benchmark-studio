"""Build the concise portfolio from its versioned Markdown source.

Run with the Python runtime returned by load_workspace_dependencies.
Render the resulting DOCX with the documents skill before delivery.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/portfolio/AI产品经理作品集.md"
OUTPUT = SOURCE.with_suffix(".docx")


def font(run, size=None, bold=None):
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string("000000")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def inline(paragraph, text, size=None):
    for part in re.split(r"(\[[^\]]+\]\([^)]+\))", text):
        match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if not match:
            font(paragraph.add_run(part), size)
            continue
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "145A86")
        properties.append(color)
        fonts = OxmlElement("w:rFonts")
        for key, value in [("ascii", "Arial"), ("hAnsi", "Arial"), ("eastAsia", "Heiti SC")]:
            fonts.set(qn("w:" + key), value)
        properties.append(fonts)
        run.append(properties)
        label = OxmlElement("w:t")
        label.text = match[1]
        run.append(label)
        link.append(run)
        paragraph._p.append(link)


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    if len(rows[0]) == 4:
        widths = [1.02, 1.62, 1.94, 2.38]
    elif len(rows[0]) == 3:
        widths = [1.88, 2.48, 2.60]
        if rows[0][-1] == "优先级":
            widths = [2.83, 3.43, 0.70]
    else:
        widths = [2.35, 4.61]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + edge)
        for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
            border.set(qn("w:" + key), value)
        borders.append(border)
    properties.append(borders)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        row_properties = table.rows[-1]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_properties = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge, amount in [("top", 75), ("bottom", 75), ("left", 90), ("right", 90)]:
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), str(amount))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tc_properties.append(margins)
            if index == 0:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "E8EFF5")
                tc_properties.append(fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            if len(value) <= 5 or (index == 0 and column == 0):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, value, 10)
            for run in paragraph.runs:
                run.bold = index == 0
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    font(spacer.add_run(""), 3)


def build():
    source = SOURCE.read_text()
    document = Document()
    document.core_properties.title = "AI测评内容生产工作台 AI产品经理作品集"
    document.core_properties.subject = "需求分析 模型策略 工作流 输入输出 节点与端到端评测"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    section = document.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = Inches(0.61)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    for name, size in [("Normal", 10.5), ("Title", 22), ("Heading 1", 17), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.18
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.keep_with_next = True
    # Remove Word template paragraph rules, including the default blue Title line.
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("AI Benchmark Studio   "), 8)
    number = OxmlElement("w:fldSimple")
    number.set(qn("w:instr"), "PAGE")
    footer._p.append(number)
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- page -->":
            document.add_page_break()
        elif line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            inline(paragraph, line[4:])
        elif line.startswith("|"):
            batch = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                batch.append(lines[index])
                index += 1
            add_table(document, batch)
            continue
        elif line.startswith("!["):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.right_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.keep_together = True
            while index < len(lines) and lines[index].strip().startswith("!["):
                match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", lines[index].strip())
                if not match:
                    raise ValueError("Invalid image syntax")
                run = paragraph.add_run()
                shape = run.add_picture(str(SOURCE.parent / match[2]), width=Inches(3.15))
                shape._inline.docPr.set("descr", match[1])
                paragraph.add_run(" ")
                index += 1
            continue
        elif line.startswith("- "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.10)
            paragraph.paragraph_format.first_line_indent = Inches(-0.10)
            inline(paragraph, "• " + line[2:])
        else:
            paragraph = document.add_paragraph()
            inline(paragraph, line)
        index += 1
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
